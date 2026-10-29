"""文件职责：解析 PDF、PPTX、DOCX 和图片，并用 Qwen-VL 抽取结构化证据。"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import mimetypes
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from merchpilot.agents.contracts import MultimodalEvidenceBatch
from merchpilot.agents.prompts import (
    MULTIMODAL_PROMPT_VERSION,
    build_multimodal_extraction_prompt,
)
from merchpilot.core.config import get_settings
from merchpilot.llm.gateway import get_multimodal_model, guarded_model_invoke
from merchpilot.services.tabular_import import EvidenceDraft

SUPPORTED_MULTIMODAL_SUFFIXES = frozenset(
    {".pdf", ".pptx", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
)
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"})
TABULAR_SUFFIXES = frozenset({".csv", ".xlsx", ".xlsm"})


@dataclass(slots=True)
class MultimodalPart:
    """一个待交给 Qwen-VL 的文件片段。"""

    locator: str
    text: str = ""
    image_bytes: bytes | None = None
    image_mime: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def source_type_for_filename(filename: str) -> str:
    """根据扩展名返回导入任务类型，供 API、Worker 和页面共用。"""
    suffix = Path(filename).suffix.lower()
    if suffix in TABULAR_SUFFIXES:
        return "tabular"
    if suffix in SUPPORTED_MULTIMODAL_SUFFIXES:
        if suffix in IMAGE_SUFFIXES:
            return "image"
        return suffix.removeprefix(".")
    raise ValueError(f"unsupported file type: {suffix or '<none>'}")


def is_supported_filename(filename: str) -> bool:
    """判断文件是否属于当前导入链路支持的格式。"""
    try:
        source_type_for_filename(filename)
    except ValueError:
        return False
    return True


def _image_part(locator: str, content: bytes, suffix: str, **metadata: Any) -> MultimodalPart:
    """构造图片片段并保留来源定位。"""
    content, suffix = _normalize_image_bytes(content, suffix)
    mime = mimetypes.guess_type(f"file{suffix}")[0] or "image/png"
    return MultimodalPart(
        locator=locator,
        image_bytes=content,
        image_mime=mime,
        metadata=metadata,
    )


def _normalize_image_bytes(content: bytes, suffix: str) -> tuple[bytes, str]:
    """把不适合直接发送给视觉模型的大图压缩为受控 JPEG。"""
    settings = get_settings()
    unsupported_direct_suffixes = {".tif", ".tiff", ".bmp"}
    if (
        len(content) <= settings.multimodal_max_image_bytes
        and suffix.lower() not in unsupported_direct_suffixes
    ):
        return content, suffix
    try:
        from PIL import Image

        with Image.open(io.BytesIO(content)) as image:
            image = image.convert("RGB")
            image.thumbnail((2400, 2400))
            output = io.BytesIO()
            image.save(output, format="JPEG", quality=82, optimize=True)
            return output.getvalue(), ".jpg"
    except Exception:
        return content, suffix


def _extract_pdf_parts(content: bytes, max_pages: int) -> list[MultimodalPart]:
    """把 PDF 每一页渲染为图片，同时保留 PDF 原生文本。"""
    import fitz

    document = fitz.open(stream=content, filetype="pdf")
    parts: list[MultimodalPart] = []
    try:
        for page_number in range(min(len(document), max_pages)):
            page = document.load_page(page_number)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            parts.append(
                _image_part(
                    f"page:{page_number + 1}",
                    pixmap.tobytes("png"),
                    ".png",
                    page_number=page_number + 1,
                    native_text=page.get_text("text")[:12_000],
                )
            )
    finally:
        document.close()
    return parts


def _extract_pptx_parts(content: bytes, max_pages: int) -> list[MultimodalPart]:
    """拆分 PPTX 幻灯片文字和嵌入图片，图片交给 Qwen-VL 识别。"""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    parts: list[MultimodalPart] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number > max_pages:
            break
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        structured_blocks = list(texts)
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = shape.table
                structured_blocks.append(
                    "表格：\n"
                    + "\n".join(
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in table.rows
                    )
                )
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                chart_rows: list[str] = []
                for series in chart.series:
                    values = [str(value) for value in series.values]
                    chart_rows.append(f"{series.name}: {', '.join(values)}")
                if chart_rows:
                    structured_blocks.append("图表数据：\n" + "\n".join(chart_rows))
        slide_text = "\n".join(structured_blocks)[:12_000]
        image_found = False
        for image_index, shape in enumerate(slide.shapes, start=1):
            if not hasattr(shape, "image"):
                continue
            image = shape.image
            parts.append(
                _image_part(
                    f"slide:{slide_number}:image:{image_index}",
                    image.blob,
                    Path(image.filename).suffix.lower() or ".png",
                    slide_number=slide_number,
                    slide_text=slide_text,
                )
            )
            image_found = True
        if not image_found and slide_text:
            parts.append(
                MultimodalPart(
                    locator=f"slide:{slide_number}",
                    text=slide_text,
                    metadata={"slide_number": slide_number},
                )
            )
    return parts


def _extract_docx_parts(content: bytes, max_pages: int) -> list[MultimodalPart]:
    """提取 DOCX 正文、表格和嵌入图片；DOCX 不强制依赖 Office 渲染器。"""
    from docx import Document

    document = Document(io.BytesIO(content))
    text_blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text_blocks.append("\n".join(row for row in rows if row.strip()))
    parts: list[MultimodalPart] = []
    full_text = "\n".join(text_blocks)[:12_000]
    if full_text:
        parts.append(MultimodalPart(locator="document:text", text=full_text))
    image_relations = [
        relation
        for relation in document.part.related_parts.values()
        if str(getattr(relation, "content_type", "")).startswith("image/")
    ]
    for image_index, relation in enumerate(image_relations, start=1):
        parts.append(
            _image_part(
                f"document:image:{image_index}",
                relation.blob,
                Path(getattr(relation, "filename", "image.png")).suffix.lower() or ".png",
                document_text=full_text,
            )
        )
        if len(parts) >= max_pages:
            break
    return parts[:max_pages]


def extract_multimodal_parts(
    content: bytes,
    filename: str,
    *,
    max_pages: int | None = None,
) -> list[MultimodalPart]:
    """把非结构化文件拆成可供 OCR/多模态抽取的片段。"""
    source_type = source_type_for_filename(filename)
    settings = get_settings()
    page_limit = max_pages or settings.multimodal_max_pages
    suffix = Path(filename).suffix.lower()
    if source_type == "image":
        return [_image_part("image:1", content, suffix, filename=filename)]
    if suffix == ".pdf":
        return _extract_pdf_parts(content, page_limit)
    if suffix == ".pptx":
        return _extract_pptx_parts(content, page_limit)
    if suffix == ".docx":
        return _extract_docx_parts(content, page_limit)
    raise ValueError(f"unsupported multimodal file type: {suffix}")


def _json_from_response(response: Any) -> dict[str, Any]:
    """解析 Qwen-VL 返回的 JSON 对象，兼容 Markdown 代码块。"""
    value = getattr(response, "content", response)
    if isinstance(value, list):
        value = "".join(
            item.get("text", "")
            for item in value
            if isinstance(item, dict) and isinstance(item.get("text"), str)
        )
    text = str(value).strip()
    if text.startswith("```"):
        text = text.removeprefix("```").removeprefix("json").removesuffix("```").strip()
    parsed = json.loads(text)
    if not isinstance(parsed, dict):
        raise ValueError("multimodal model response must be a JSON object")
    return parsed


def _build_prompt(part: MultimodalPart, filename: str) -> str:
    """构造 OCR 与商品证据抽取 Prompt，明确事实、推断和待审核边界。"""
    auxiliary_text = (
        part.text
        or part.metadata.get("native_text")
        or part.metadata.get("slide_text")
        or part.metadata.get("document_text")
        or "无"
    )
    return build_multimodal_extraction_prompt(filename, part.locator, auxiliary_text)


async def _extract_part(model: Any, part: MultimodalPart, filename: str) -> MultimodalEvidenceBatch:
    """调用 Qwen-VL 抽取单个文件片段。"""
    prompt = _build_prompt(part, filename)
    if part.image_bytes:
        encoded = base64.b64encode(part.image_bytes).decode("ascii")
        image_url = f"data:{part.image_mime or 'image/png'};base64,{encoded}"
        message: Any = [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": image_url}},
                ],
            }
        ]
    else:
        message = prompt
    response = await guarded_model_invoke(model, message)
    return MultimodalEvidenceBatch.model_validate(_json_from_response(response))


def _fallback_item(part: MultimodalPart, filename: str) -> dict[str, Any] | None:
    """无模型时保留可检索的原生文本，不把图片内容伪装成已识别事实。"""
    text = (part.text or part.metadata.get("native_text") or part.metadata.get("slide_text") or "").strip()
    if not text:
        return None
    return {
        "title": f"{filename} {part.locator}",
        "content": text[:12_000],
        "data_type": "document_text_fallback",
        "evidence_tags": ["multimodal_fallback"],
        "confidence": "low",
        "needs_review": True,
    }


def _parse_observed_at(value: object) -> datetime | None:
    """把模型识别出的日期转换为统一 UTC 时间，无法确认时保持空值。"""
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


async def extract_multimodal_evidence(
    content: bytes,
    filename: str,
    *,
    import_job_id: str,
) -> tuple[list[EvidenceDraft], dict[str, Any]]:
    """抽取非结构化文件并转换为统一 EvidenceDraft。"""
    parts = extract_multimodal_parts(content, filename)
    model_error: str | None = None
    fallback_count = 0
    try:
        model = get_multimodal_model()
    except Exception as exc:
        model = None
        model_error = str(exc)

    extracted: list[tuple[MultimodalPart, dict[str, Any]]] = []
    for part in parts:
        if model is None:
            item = _fallback_item(part, filename)
            if item:
                extracted.append((part, item))
            continue
        try:
            batch = await _extract_part(model, part, filename)
            extracted.extend((part, item.model_dump(mode="json")) for item in batch.items)
        except Exception as exc:
            model_error = str(exc)
            fallback_count += 1
            item = _fallback_item(part, filename)
            if item:
                extracted.append((part, item))

    drafts: list[EvidenceDraft] = []
    for index, (part, item) in enumerate(extracted, start=1):
        text = str(item.get("content") or "").strip()
        if not text:
            continue
        source_id = f"{filename}:{part.locator}:{index}"
        content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
        evidence_key = f"{import_job_id}:{source_id}:{content_hash}".encode()
        evidence_id = f"evidence_{hashlib.sha256(evidence_key).hexdigest()[:32]}"
        needs_review = bool(item.get("needs_review", True))
        drafts.append(
            EvidenceDraft(
                id=evidence_id,
                source_id=source_id,
                source_name=filename,
                source_row=index,
                content=text,
                metadata={
                    "import_job_id": import_job_id,
                    "source_locator": part.locator,
                    "source_format": Path(filename).suffix.lower().removeprefix("."),
                    "extraction": (
                        "qwen_vl"
                        if model is not None and fallback_count == 0
                        else "qwen_vl_partial"
                        if model is not None
                        else "native_text_fallback"
                    ),
                    "extraction_model": getattr(model, "model_name", None),
                    "prompt_version": MULTIMODAL_PROMPT_VERSION,
                    "title": item.get("title"),
                    "category": item.get("category"),
                    "target_market": item.get("target_market"),
                    "target_user": item.get("target_user"),
                    "data_type": item.get("data_type") or "document_extract",
                    "observed_at": item.get("observed_at"),
                    "evidence_tags": item.get("evidence_tags", []),
                    "confidence": item.get("confidence", "low"),
                    "needs_review": needs_review,
                    "extraction_error": model_error,
                },
                content_hash=content_hash,
                quality_status="warning" if needs_review else "verified",
                observed_at=_parse_observed_at(item.get("observed_at")),
            )
        )
    return drafts, {
        "parts": len(parts),
        "extracted_items": len(drafts),
        "extraction": (
            "qwen_vl"
            if model is not None and fallback_count == 0
            else "qwen_vl_partial"
            if model is not None
            else "native_text_fallback"
        ),
        "model": getattr(model, "model_name", None),
        "prompt_version": MULTIMODAL_PROMPT_VERSION,
        "error": model_error,
    }
