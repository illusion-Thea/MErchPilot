"""文件职责：验证 import quality 相关规则、边界与回归行为。"""

from io import BytesIO

from openpyxl import load_workbook

from merchpilot.infrastructure.models import DataImportJob, EvidenceRecord
from merchpilot.services.import_quality import build_import_quality_report_bytes


def test_quality_report_contains_issues_and_sanitizes_external_text() -> None:
    """验证该测试函数描述的业务场景。

    确保业务规则、权限与异常边界保持稳定。
    """
    job = DataImportJob(
        id="import-1",
        tenant_id="default",
        filename="collection.xlsx",
        storage_key="imports/collection.xlsx",
        status="completed",
        total_rows=1,
        valid_rows=0,
        invalid_rows=1,
        created_by="local-user",
    )
    record = EvidenceRecord(
        id="evidence-1",
        tenant_id="default",
        import_job_id=job.id,
        source_type="tabular",
        source_name=job.filename,
        source_row=2,
        source_id="=CMD()",
        content='=HYPERLINK("https://example.com")',
        metadata_json={
            "normalized_row": {
                "data_type": "用户评论",
                "category": "智能保温杯",
                "target_market": "中国",
                "data_source": "公开网页",
            },
            "quality_issues": [
                {
                    "field": "source_url",
                    "code": "missing_source_url",
                    "message": "公开采集缺少来源URL",
                    "severity": "warning",
                }
            ],
        },
        content_hash="hash-1",
        quality_status="blocked",
        index_status="failed",
    )

    workbook = load_workbook(BytesIO(build_import_quality_report_bytes(job, [record])))

    assert workbook.sheetnames == ["质量汇总", "问题明细"]
    detail_sheet = workbook["问题明细"]
    assert detail_sheet["B2"].value == "'=CMD()"
    assert detail_sheet["H2"].value == "公开采集缺少来源URL"
    assert detail_sheet["I2"].value == '\'=HYPERLINK("https://example.com")'
