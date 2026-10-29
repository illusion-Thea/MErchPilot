"""文件职责：验证非结构化文件解析、Qwen-VL 消息和无模型降级。"""

from __future__ import annotations

import io
from types import SimpleNamespace

import fitz
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from merchpilot.services.multimodal_import import (
    extract_multimodal_evidence,
    extract_multimodal_parts,
    source_type_for_filename,
)


def _png_bytes() -> bytes:
    """生成最小图片样本，避免测试依赖仓库外部素材。"""
    output = io.BytesIO()
    Image.new("RGB", (32, 32), color="white").save(output, format="PNG")
    return output.getvalue()


def test_source_type_supports_tabular_and_multimodal_files() -> None:
    """所有目标格式都必须映射到统一导入类型。"""
    assert source_type_for_filename("data.xlsx") == "tabular"
    assert source_type_for_filename("report.pdf") == "pdf"
    assert source_type_for_filename("deck.pptx") == "pptx"
    assert source_type_for_filename("brief.docx") == "docx"
    assert source_type_for_filename("photo.png") == "image"


def test_extract_pdf_pptx_docx_and_image_parts() -> None:
    """PDF、PPTX、DOCX 和图片都能拆成可追溯片段。"""
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "智能水杯 市场价格 199")
    pdf_content = pdf.tobytes()
    pdf.close()

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "通勤用户 防漏"
    pptx_output = io.BytesIO()
    presentation.save(pptx_output)

    document = Document()
    document.add_paragraph("用户评论：杯盖清洗方便")
    docx_output = io.BytesIO()
    document.save(docx_output)

    assert extract_multimodal_parts(pdf_content, "report.pdf")[0].locator == "page:1"
    assert extract_multimodal_parts(pptx_output.getvalue(), "deck.pptx")[0].locator == "slide:1"
    assert extract_multimodal_parts(docx_output.getvalue(), "brief.docx")[0].text
    assert extract_multimodal_parts(_png_bytes(), "photo.png")[0].image_bytes


@pytest.mark.asyncio
async def test_multimodal_import_falls_back_to_native_text_without_qwen(monkeypatch) -> None:
    """没有 DashScope Key 时，文档原生文本仍能进入证据链并标记待审核。"""
    def unavailable_model():
        """模拟本地未配置 Qwen-VL。"""
        raise RuntimeError("DASHSCOPE_API_KEY is not configured")

    monkeypatch.setattr(
        "merchpilot.services.multimodal_import.get_multimodal_model",
        unavailable_model,
    )
    document = Document()
    document.add_paragraph("供应商交期为 15 天，MOQ 为 500 件")
    content = io.BytesIO()
    document.save(content)

    drafts, summary = await extract_multimodal_evidence(
        content.getvalue(),
        "supplier.docx",
        import_job_id="import-test",
    )

    assert summary["extraction"] == "native_text_fallback"
    assert drafts
    assert drafts[0].metadata["needs_review"] is True
    assert drafts[0].quality_status == "warning"
    assert "MOQ" in drafts[0].content


@pytest.mark.asyncio
async def test_qwen_vl_message_contains_standard_image_url(monkeypatch) -> None:
    """Qwen-VL 请求必须使用 OpenAI 兼容的 role/content/image_url 结构。"""
    captured: dict[str, object] = {}

    class StubModel:
        """捕获视觉模型输入并返回最小合法抽取结果。"""

        model_name = "qwen-vl-max"

        async def ainvoke(self, message: object):
            captured["message"] = message
            return SimpleNamespace(
                content='{"items":[{"content":"图中显示价格199元","confidence":"high","needs_review":false}]}'
            )

    monkeypatch.setattr(
        "merchpilot.services.multimodal_import.get_multimodal_model",
        lambda: StubModel(),
    )
    drafts, summary = await extract_multimodal_evidence(
        _png_bytes(),
        "price.png",
        import_job_id="import-vl",
    )

    message = captured["message"]
    assert isinstance(message, list)
    assert message[0]["role"] == "user"
    assert message[0]["content"][1]["type"] == "image_url"
    assert summary["extraction"] == "qwen_vl"
    assert drafts[0].quality_status == "verified"
