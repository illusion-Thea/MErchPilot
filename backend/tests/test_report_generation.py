"""文件职责：验证 report generation 相关规则、边界与回归行为。"""

from io import BytesIO

from openpyxl import load_workbook

from merchpilot.infrastructure.models import PlanningProject, PlanningRun
from merchpilot.services.report_generation import build_planning_report_bytes


def test_build_planning_report_contains_business_sheets_and_sanitizes_formulas() -> None:
    """验证该测试函数描述的业务场景。

    确保业务规则、权限与异常边界保持稳定。
    """
    project = PlanningProject(
        id="project-1",
        tenant_id="default",
        name="智能水杯规划",
        category="智能水杯",
        target_market="中国",
        status="approval_required",
        constraints={"target_user": "=HYPERLINK(\"https://example.com\")"},
        created_by="local-user",
    )
    run = PlanningRun(
        id="run-1",
        project_id=project.id,
        thread_id="thread-1",
        status="approval_required",
        input_snapshot={},
        result_snapshot={
            "outputs": {
                "scorecard": {
                    "score": 82,
                    "status": "ready",
                    "evidence_count": 1,
                    "breakdown": {"evidence": {"score": 30, "max_score": 30}},
                },
                "finance": {"margin_rate": 0.42, "selling_price": 99},
                "supplier": {"risk_level": "low"},
                "concept": {
                    "name": "智能水杯轻量通勤款",
                    "selected_candidate_id": "concept-1",
                    "selected_candidate": {"name": "智能水杯轻量通勤款"},
                    "candidates": [
                        {
                            "id": "concept-1",
                            "name": "智能水杯轻量通勤款",
                            "recommendation": "优先验证",
                            "scorecard": {"total_score": 77},
                            "financial_scenario": {"margin_rate": 0.5, "estimated_profit": 1000},
                            "validation_plan": {
                                "status": "partial",
                                "status_label": "部分具备，需补齐验证",
                                "completion_rate": 75,
                                "next_actions": ["完成样品测试"],
                                "gates": [
                                    {
                                        "id": "specification",
                                        "label": "规格与样品",
                                        "materials": [{"filename": "样品报告.pdf", "note": "测试通过"}],
                                    }
                                ],
                            },
                            "positioning": "便携通勤",
                            "target_user": "城市白领",
                            "evidence_features": ["防漏"],
                            "pain_points_addressed": ["难清洗"],
                            "feature_hypotheses": ["轻量化"],
                            "price_reference": {"suggested_price": 129},
                            "unit_cost_reference": 50,
                            "differentiation_hypothesis": "轻量",
                            "confidence": "medium",
                            "risks": ["需验证"],
                            "evidence_ids": ["evidence-1"],
                        }
                    ],
                },
                "market": {
                    "hits": [
                        {
                            "id": "evidence-1",
                            "source_name": "市场调研",
                            "score": 0.9,
                            "content": "=CMD()",
                        }
                    ]
                },
            }
        },
    )

    workbook = load_workbook(BytesIO(build_planning_report_bytes(project, run)))

    assert workbook.sheetnames == ["方案摘要", "评分拆解", "约束参数", "候选商品", "规划建议", "证据明细"]
    assert workbook["约束参数"]["B2"].value == "'=HYPERLINK(\"https://example.com\")"
    assert workbook["证据明细"]["E2"].value == "'=CMD()"
    assert workbook["候选商品"]["B2"].value == "已选"
    assert workbook["候选商品"]["C2"].value == "智能水杯轻量通勤款"
    assert workbook["候选商品"]["D2"].value == 77
    assert workbook["候选商品"]["E2"].value == "优先验证"
    assert workbook["候选商品"]["F2"].value == "部分具备，需补齐验证"
    assert workbook["候选商品"]["G2"].value == 75
    assert workbook["候选商品"]["H2"].value == "完成样品测试"
    assert workbook["候选商品"]["I2"].value == "50.00%"
    assert workbook["候选商品"]["J2"].value == 1000
    assert workbook["候选商品"]["V2"].value == "规格与样品: 样品报告.pdf（测试通过）"
